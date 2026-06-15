import os
from functools import wraps
from dotenv import load_dotenv
load_dotenv()

API_KEY = os.getenv("API_KEY")

from db.models import Base, engine, User, Device, NetworkLink, PingHistory, Alert, Category
from monitor.monitor import check_device, monitor_status
from db.database import (
    SessionLocal, login, add_device, add_devices_bulk,
    get_topology, get_locations, get_location_devices, get_location_links,
    get_device_by_ip, get_links_for_device, get_devices_stats,
    all_devices_status, get_details_paginated,
    update_device, delete_device,
    get_all_links, add_link, update_link, delete_link,
    get_links_paginated, get_link_filter_options, category_list, category_add, 
    get_dashboard_state, save_dashboard_state,
    get_cached_topology, cache_topology, invalidate_cache_for,
    create_user, get_all_users, delete_user, reset_user_password
)
from flask import Flask, request, redirect, flash, render_template, session, jsonify, Response, send_file
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from flask_limiter.errors import RateLimitExceeded
from flask_login import LoginManager, login_required, login_user, logout_user, current_user
from extensions import socketio
from monitor.scheduler import scheduler
from werkzeug.security import generate_password_hash
from captcha.image import ImageCaptcha
import random
import string
import io
import threading
import secrets
import time
from datetime import datetime, timezone
import ipaddress

def is_valid_ip(ip):
    """
    Validate IPv4/IPv6 address.

    Returns:
        (True, None) if valid
        (False, error_message) if invalid
    """
    if not ip:
        return False, "IP address is required"

    try:
        ipaddress.ip_address(ip.strip())
        return True, None
    except ValueError:
        return False, f"'{ip}' is not a valid IP address"

def generate_captcha(length=6):
    return ''.join(random.choices(string.ascii_uppercase + string.digits, k=length))


app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "change_this_to_random_secret")
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    PERMANENT_SESSION_LIFETIME=3600,
    SESSION_COOKIE_NAME="session",
)
socketio.init_app(app)

limiter = Limiter(
    get_remote_address,
    app=app,
    default_limits=["200 per day", "50 per hour"],
    storage_uri="memory://",
)

@app.errorhandler(RateLimitExceeded)
def handle_rate_limit(e):
    flash("Too many login requests. Please try again later.")
    return redirect("/")


@limiter.request_filter
def exempt_api():
    return request.path.startswith("/api/")


def require_api_key(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        api_key = request.headers.get("X-API-Key")
        csrf_token = request.headers.get("X-CSRF-Token")
        if API_KEY and api_key == API_KEY:
            return f(*args, **kwargs)
        if csrf_token and csrf_token == session.get("csrf_token"):
            return f(*args, **kwargs)
        return jsonify({"error": "Unauthorized"}), 401
    return decorated

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "/"


@app.context_processor
def inject_csrf_token():
    return dict(csrf_token=session.get("csrf_token", ""))


@login_manager.user_loader
def load_user(user_id):
    s = SessionLocal()
    try:
        return s.get(User, int(user_id))
    finally:
        s.close()


@app.before_request
def session_timeout_check():
    if current_user.is_authenticated:
        now = time.time()
        last_activity = session.get("last_activity", now)
        if now - last_activity > 3600:
            logout_user()
            session.clear()
            flash("Session expired. Please log in again.")
            return redirect("/")
        session["last_activity"] = now
        session.permanent = True


@app.before_request
def force_https():
    if not app.debug and os.getenv("SSL_CERTFILE") and os.getenv("SSL_KEYFILE"):
        if not request.is_secure and request.headers.get("X-Forwarded-Proto", "http") != "https":
            url = request.url.replace("http://", "https://", 1)
            return redirect(url, code=301)


@app.before_request
def check_must_change_password():
    if current_user.is_authenticated and current_user.must_change_password:
        if request.endpoint not in ("change_password", "index", "captcha_image", "captcha_refresh", "logout", "static"):
            return redirect("/change-password")


@app.after_request
def add_security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["X-XSS-Protection"] = "1; mode=block"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    if request.path.startswith("/api/") or request.path.endswith(".html"):
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
    return response


def create_default_admin():
    s = SessionLocal()
    try:
        if not s.query(User).first():
            admin_user = os.getenv("DEFAULT_ADMIN_USER", "admin")
            admin_pass = os.getenv("DEFAULT_ADMIN_PASS", "admin")
            admin = User(
                user=admin_user,
                passwd=generate_password_hash(admin_pass),
                must_change_password=False
            )
            s.add(admin)
            s.commit()
            print(f"Default admin user created ({admin_user}/{admin_pass})")
    finally:
        s.close()


@app.route("/captcha-image")
def captcha_image():
    text = session.get("captcha", "ABCD")
    image = ImageCaptcha(width=200, height=60)
    buf = io.BytesIO()
    image.write(text, buf)
    buf.seek(0)
    return Response(buf.getvalue(), mimetype="image/png")


@app.route("/captcha-refresh")
def captcha_refresh():
    session["captcha"] = generate_captcha()
    session.modified = True
    image = ImageCaptcha(width=200, height=60)
    buf = io.BytesIO()
    image.write(session["captcha"], buf)
    buf.seek(0)
    return Response(buf.getvalue(), mimetype="image/png")


@app.route("/", methods=["GET", "POST"])
@limiter.limit("10 per minute", methods=["POST"])
def index():
    if request.method == "POST":
        user = request.form.get("user")
        passwd = request.form.get("passwd")
        captcha = request.form.get("captcha").upper()
        print(captcha)
        if captcha != session.get("captcha"):
            flash("Invalid Captcha")
            return redirect("/")
        response = login(username=user, password=passwd)
        if response is False:
            flash("Invalid Credentials")
        elif isinstance(response, Exception):
            flash(f"Error: {response}")
        else:
            login_user(response)
            session["csrf_token"] = secrets.token_hex(32)
            if response.must_change_password:
                return redirect("/change-password")
            return redirect("/dashboard")
        return render_template("index.html")
    if request.method=="GET" and current_user.is_authenticated:
        return redirect("/dashboard")
    session["captcha"] = generate_captcha()
    return render_template("index.html")


@app.route("/change-password", methods=["GET", "POST"])
@login_required
def change_password():
    if request.method == "POST":
        current_pw = request.form.get("current_password")
        new_pw = request.form.get("new_password")
        confirm_pw = request.form.get("confirm_password")
        from werkzeug.security import check_password_hash
        if not check_password_hash(current_user.passwd, current_pw):
            flash("Current password is incorrect")
            return redirect("/change-password")
        if new_pw != confirm_pw:
            flash("New passwords do not match")
            return redirect("/change-password")
        if len(new_pw) < 4:
            flash("Password must be at least 4 characters")
            return redirect("/change-password")
        s = SessionLocal()
        try:
            user = s.get(User, current_user.id)
            user.passwd = generate_password_hash(new_pw)
            user.must_change_password = False
            s.commit()
            flash("Password changed successfully")
        except Exception as e:
            s.rollback()
            flash(f"Error: {e}")
        finally:
            s.close()
        return redirect("/dashboard")
    return render_template("change_password.html")

@limiter.exempt()
@app.route("/dashboard")
@login_required
def dashboard():
    page         = request.args.get("page", 1, type=int)
    per_page     = request.args.get("per_page", 20, type=int)
    search_ip    = request.args.get("search_ip", "").strip()
    filter_status    = request.args.get("filter_status", "").strip()
    filter_category  = request.args.get("filter_category", "").strip()
    filter_location  = request.args.get("filter_location", "").strip()
    sort_by      = request.args.get("sort_by", "id").strip()
    sort_order   = request.args.get("sort_order", "desc").strip()
    edit_id      = request.args.get("edit", type=int)

    if per_page not in (10, 20, 50, 100):
        per_page = 20
    page = max(1, page)

    valid_sort_cols = ("id", "source_ip", "destination_ip", "source_location",
                       "destination_location", "remark", "status")
    if sort_by not in valid_sort_cols:
        sort_by = "id"
    if sort_order not in ("asc", "desc"):
        sort_order = "desc"

    links, total, total_pages = get_links_paginated(
        page=page, per_page=per_page,
        search_ip=search_ip,
        filter_status=filter_status,
        filter_category=filter_category,
        filter_location=filter_location,
        sort_by=sort_by,
        sort_order=sort_order,
    )

    edit_link = None
    if edit_id:
        all_lks = get_all_links()
        edit_link = next((l for l in all_lks if l["id"] == edit_id), None)

    categories, locations = get_link_filter_options()
    need_location_ip = session.get("need_location_ip", None)
    pending_link = session.get("pending_link", None)

    return render_template(
        "radio3.html",
        links=links,
        edit_link=edit_link,
        active_page='dashboard',
        page=page,
        per_page=per_page,
        total=total,
        total_pages=total_pages,
        search_ip=search_ip,
        filter_status=filter_status,
        filter_category=filter_category,
        filter_location=filter_location,
        sort_by=sort_by,
        sort_order=sort_order,
        categories=categories,
        locations=locations,
        need_location_ip=need_location_ip,
        pending_link=pending_link,
    )


@app.route("/dashboard/radio/<id>", methods=['GET', 'POST'])
@login_required
def radio():
    if request.method == "GET":
        return render_template("detai.html", active_page='dashboard')
    if request.method == "POST":
        ip = request.form.get("ip")
        if ip:
            return redirect(f"/device/{ip}")
        return redirect("/dashboard")


@app.route("/devices/upload/template")
@login_required
def upload_devices_template():
    import openpyxl
    from io import BytesIO
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Device Template"
    headers = ["Source IP", "Source Model", "Source Category", "Source Location","Destination IP", "Destination Model", "Destination Category", "Destination Location","Remark"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 25
        cell = ws.cell(row=1, column=col)
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill(start_color="0f172a", end_color="0f172a", fill_type="solid")
        cell.alignment = openpyxl.styles.Alignment(horizontal="center")
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, download_name="device_upload_template.xlsx", as_attachment=True, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/devices/upload", methods=["GET", "POST"])
@login_required
def upload_devices():
    if request.method == "POST":
        file = request.files.get("file")

        if not file or file.filename == "":
            flash("No file selected")
            return redirect("/devices/upload")

        try:
            import openpyxl

            wb = openpyxl.load_workbook(file)
            ws = wb.active

            device_categories = set(category_list())

            devices_data = []
            links_data = []

            # Skip header row
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 1:
                    continue

                (
                    source_ip,
                    source_model,
                    source_category,
                    source_location,
                    destination_ip,
                    destination_model,
                    destination_category,
                    destination_location,
                    remark
                ) = (row + (None,)*9)[:9]

                # Source device
                if source_ip:
                    source_ip = str(source_ip).strip()

                    if not source_category:
                        flash(f"Missing source category for {source_ip}")
                        continue

                    src_cat = str(source_category).strip().lower()

                    if src_cat not in device_categories:
                        category_add(src_cat)
                        device_categories.add(src_cat)

                    devices_data.append({
                        "device_name": source_ip,
                        "ip": source_ip,
                        "model": str(source_model).strip() if source_model else "",
                        "category": src_cat,
                        "location": str(source_location).strip() if source_location else "",
                    })

                # Destination device
                if destination_ip:
                    destination_ip = str(destination_ip).strip()

                    if not destination_category:
                        flash(f"Missing destination category for {destination_ip}")
                        continue

                    dst_cat = str(destination_category).strip().lower()

                    if dst_cat not in device_categories:
                        category_add(dst_cat)
                        device_categories.add(dst_cat)

                    devices_data.append({
                        "device_name": destination_ip,
                        "ip": destination_ip,
                        "model": str(destination_model).strip() if destination_model else "",
                        "category": dst_cat,
                        "location": str(destination_location).strip() if destination_location else "",
                    })

                # Network link
                if source_ip and destination_ip:
                    links_data.append({
                        "source_ip": source_ip,
                        "destination_ip": destination_ip,
                        "remark":remark
                    })
                if source_ip and not destination_ip:
                    links_data.append({
                        "source_ip": source_ip,
                        "destination_ip": None,
                        "remark":remark
                    })



            if devices_data:

                added_devices, device_errors = add_devices_bulk(devices_data)

                added_links = 0
                link_errors = []

                for idx, link in enumerate(links_data, start=1):
                    try:
                        add_link(
                            source_ip=link["source_ip"],
                            destination_ip=link["destination_ip"],
                            remark=link.get("remark")
                        )
                        added_links += 1
                    except KeyError as e:
                        link_errors.append(f"Row {idx}: Missing field {e}")
                    except ValueError as e:
                        link_errors.append(f"Row {idx}: Invalid data - {e}")
                    except Exception as e:
                        app.logger.exception("Failed to add link")
                        link_errors.append(f"Row {idx}: Unexpected error - {e}")

                flash(f"Added {len(added_devices)} devices")
                flash(f"Added {added_links} network links")

                for err in device_errors:
                    flash(f"Device Error: {err}")

                for err in link_errors:
                    flash(f"Link Error: {err}")

            else:
                flash("No valid devices found in file")

        except Exception as e:
            flash(f"Error processing file: {e}")

        return redirect("/devices/upload")

    return render_template("upload_devices.html",active_page="upload")


@app.route("/location/<location_name>")
@login_required
def location_page(location_name):
    devices = get_location_devices(location_name)
    links = get_location_links(location_name)
    total = len(devices)
    down = sum(1 for d in devices if not d["status"])
    return render_template("location.html",
                           location_name=location_name,
                           devices=devices,
                           links=links,
                           total_devices=total,
                           down_devices=down,
                           active_page='locations')


@app.route("/devices", methods=["GET", "POST"])
@login_required
def devices():
    edit_ip = request.args.get("edit")
    edit = get_device_by_ip(edit_ip) if edit_ip else None

    if request.method == "POST":
        device_name = request.form.get("device_name", "").strip()
        ip = request.form.get("ip", "").strip()
        model = request.form.get("model", "").strip()
        category = request.form.get("category", "").strip()
        location = request.form.get("location", "").strip()
        existing = get_device_by_ip(ip)

        if edit:
            update_device(ip, device_name=device_name, model=model, category=category, location=location)
            invalidate_cache_for("topology")
            flash("Device updated")
        elif existing:
            flash("Device with this IP already exists")
        else:
            add_device(device_name=device_name or ip, ip=ip, model=model, category=category, location=location)
            invalidate_cache_for("topology")
            flash("Device added")
        return redirect("/devices")

    page        = request.args.get("page", 1, type=int)
    per_page    = request.args.get("per_page", 20, type=int)
    search      = request.args.get("search", "").strip()
    sort_by     = request.args.get("sort_by", "ip").strip()
    sort_order  = request.args.get("sort_order", "asc").strip()

    if per_page not in (10, 20, 50, 100):
        per_page = 20
    page = max(1, page)

    valid_sort_cols = ("ip", "device_name", "model", "category", "location", "status")
    if sort_by not in valid_sort_cols:
        sort_by = "ip"
    if sort_order not in ("asc", "desc"):
        sort_order = "asc"

    devices_list, total = get_details_paginated(
        page=page, per_page=per_page,
        sort_by=sort_by, sort_dir=sort_order, search=search,
    )
    total_pages = max(1, (total + per_page - 1) // per_page)
    categories = category_list()

    return render_template("devices.html",
                           devices=devices_list, edit=edit, categories=categories,
                           active_page="devices",
                           page=page, per_page=per_page, total=total, total_pages=total_pages,
                           search=search, sort_by=sort_by, sort_order=sort_order)


@app.route("/devices/delete/<ip>")
@login_required
def delete_device_page(ip):
    result = delete_device(ip)
    if isinstance(result, Exception):
        flash(f"Error deleting device: {result}")
    elif result is False:
        flash("Device not found")
    else:
        invalidate_cache_for("topology")
        flash("Device deleted")
    return redirect("/devices")


@app.route("/device/<ip>")
@login_required
def device_page(ip):
    device = get_device_by_ip(ip)
    if not device:
        return "Device not found", 404
    links = get_links_for_device(ip)
    return render_template("device.html",
                           device=device,
                           links=links,
                           active_page='devices')


# --- Links Management ---
@limiter.exempt()
@app.route("/links", methods=["GET", "POST"])
@login_required
def links_route():
    # ── Hard cancel via GET param — most reliable escape hatch ───
    if request.args.get("cancel_pending") == "1":
        session.pop("pending_link", None)
        session.pop("need_location_ip", None)
        session.modified = True
        flash("Link addition cancelled.")
        return redirect("/links")

    # ── Stale session guard — clear if IP no longer makes sense ──
    if session.get("need_location_ip") and request.method == "GET":
        pending = session.get("pending_link")
        if not pending:
            session.pop("need_location_ip", None)
            session.modified = True

    edit_id = request.args.get("edit", type=int)
    delete_id = request.args.get("delete", type=int)

    if delete_id:
        result = delete_link(delete_id)
        if isinstance(result, Exception):
            flash(f"Error deleting link: {result}")
        else:
            invalidate_cache_for("topology")
            flash("Link deleted")
        return redirect("/links")

    if request.method == "POST":
        # ── Handle "set location" step for a pending new device ──
        if request.form.get("_action") == "set_location":
            ip_to_set    = request.form.get("_pending_ip", "").strip()
            location_val = request.form.get("_location", "").strip()
            hostname_val = request.form.get("_hostname", "").strip()
            category_val = request.form.get("_category", "").strip()
            model_val    = request.form.get("_model", "").strip()

            pending = session.get("pending_link")
            if not pending or not ip_to_set or not location_val or not category_val:
                session.pop("pending_link", None)
                session.pop("need_location_ip", None)
                session.modified = True
                if not category_val:
                    flash("Category is required.")
                else:
                    flash("Invalid state — please try adding the link again.")
                return redirect("/links")

            src = pending["source_ip"]
            dst = pending.get("destination_ip")
            is_src_step = (ip_to_set == src)

            # Create this device now (first time it's touched)
            if not get_device_by_ip(ip_to_set):
                add_device(
                    device_name=hostname_val or ip_to_set,
                    ip=ip_to_set,
                    model=model_val,
                    category=category_val or "radio",
                    location=location_val
                )
            else:
                update_device(ip_to_set,
                              device_name=hostname_val or None,
                              model=model_val or None,
                              category=category_val or None,
                              location=location_val)

            if is_src_step and dst and not get_device_by_ip(dst):
                # Source is done, destination is also new — ask for dst location next
                session["need_location_ip"] = dst
                session.modified = True
                flash(f"Device {dst} is new — please set its location to continue.")
                return redirect("/links")

            # Both devices exist now — create the link
            # Update the other device with pending data if it already existed
            other_ip = dst if is_src_step else src
            other_pending_prefix = "destination" if is_src_step else "source"
            if other_ip and get_device_by_ip(other_ip):
                update_device(other_ip,
                              device_name=pending.get(f"{other_pending_prefix}_hostname") or None,
                              model=pending.get(f"{other_pending_prefix}_model") or None,
                              category=pending.get(f"{other_pending_prefix}_category") or None,
                              location=pending.get(f"{other_pending_prefix}_location") or None)

            result = add_link(
                source_ip=src,
                destination_ip=dst or None,
                remark=pending.get("remark"),
                link_type=pending.get("link_type"),
                bandwidth=pending.get("bandwidth", ""),
            )
            if isinstance(result, Exception):
                flash(f"Error adding link: {result}")
            else:
                invalidate_cache_for("topology")
                flash(f"Link {src} → {dst or 'Unlinked'} added.")

            session.pop("pending_link", None)
            session.pop("need_location_ip", None)
            session.modified = True
            return redirect("/links")

        # ── Cancel a pending link addition ──
        if request.form.get("_action") == "cancel":
            session.pop("pending_link", None)
            session.pop("need_location_ip", None)
            session.modified = True
            flash("Link addition cancelled.")
            return redirect("/links")

        # ── Normal add / edit flow ────────────────────────────────
        source_ip            = request.form.get("source_ip", "").strip()
        source_hostname      = request.form.get("source_hostname", "").strip()
        source_location      = request.form.get("source_location", "").strip()
        source_model         = request.form.get("source_model", "").strip()
        source_category      = request.form.get("source_category", "").strip()

        destination_ip            = request.form.get("destination_ip", "").strip()
        destination_hostname      = request.form.get("destination_hostname", "").strip()
        destination_location      = request.form.get("destination_location", "").strip()
        destination_model         = request.form.get("destination_model", "").strip()
        destination_category      = request.form.get("destination_category", "").strip()

        link_type      = request.form.get("link_type", "").strip()
        remark         = request.form.get("remark", "").strip()
        link_id        = request.form.get("id", type=int)

        # Validate source IP
        valid_src, err_src = is_valid_ip(source_ip)
        if not valid_src:
            flash(f"Source IP: {err_src}")
            return redirect("/links")

        # Validate source category
        if not source_category:
            flash("Source category is required")
            return redirect("/links")

        # Validate destination IP (optional)
        if destination_ip:
            valid_dst, err_dst = is_valid_ip(destination_ip)
            if not valid_dst:
                flash(f"Destination IP: {err_dst}")
                return redirect("/links")

            # Validate destination category when destination IP is provided
            if not destination_category:
                flash("Destination category is required")
                return redirect("/links")

        if destination_ip and source_ip == destination_ip:
            flash("Source and destination IP cannot be the same")
            return redirect("/links")

        existing_src = get_device_by_ip(source_ip)
        existing_dst = get_device_by_ip(destination_ip) if destination_ip else None

        if link_id:
            # Edit existing link — auto-create or update devices
            if not existing_src:
                add_device(device_name=source_hostname or source_ip, ip=source_ip, model=source_model,
                           category=source_category or "radio", location=source_location)
            else:
                update_device(source_ip, device_name=source_hostname or None,
                              model=source_model or None,
                              category=source_category or None, location=source_location or None)

            if destination_ip:
                if not existing_dst:
                    add_device(device_name=destination_hostname or destination_ip, ip=destination_ip, model=destination_model,
                               category=destination_category or "radio", location=destination_location)
                else:
                    update_device(destination_ip, device_name=destination_hostname or None,
                                  model=destination_model or None,
                                  category=destination_category or None, location=destination_location or None)

            result = update_link(link_id, source_ip=source_ip, destination_ip=destination_ip,
                                 link_type=link_type, remark=remark)
            if isinstance(result, Exception):
                flash(f"Error updating link: {result}")
            else:
                invalidate_cache_for("topology")
                flash("Link updated")
            return redirect("/links")

        # New link — check for unknown IPs and ask for location
        missing_src = not existing_src
        missing_dst = destination_ip and not existing_dst

        if missing_src or missing_dst:
            # Stash the intended link so we can complete it after locations are set
            session["pending_link"] = {
                k: v for k, v in {
                    "source_ip":            source_ip,
                    "source_hostname":      source_hostname,
                    "source_location":      source_location,
                    "source_model":         source_model,
                    "source_category":      source_category,
                    "destination_ip":       destination_ip,
                    "destination_hostname": destination_hostname,
                    "destination_location": destination_location,
                    "destination_model":    destination_model,
                    "destination_category": destination_category,
                    "link_type":            link_type,
                    "remark":               remark,
                }.items() if v
            }
            if missing_src:
                session["need_location_ip"] = source_ip
                session.modified = True
                flash(f"Device {source_ip} is new — please set its location to continue.")
                return redirect("/links")
            else:
                session["need_location_ip"] = destination_ip
                session.modified = True
                flash(f"Device {destination_ip} is new — please set its location to continue.")
                return redirect("/links")

        # Both IPs already known — update devices then create link
        if existing_src:
            update_device(source_ip, device_name=source_hostname or None,
                          model=source_model or None,
                          category=source_category or None, location=source_location or None)
        else:
            add_device(device_name=source_hostname or source_ip, ip=source_ip, model=source_model,
                       category=source_category or "radio", location=source_location)

        if destination_ip:
            if existing_dst:
                update_device(destination_ip, device_name=destination_hostname or None,
                              model=destination_model or None,
                              category=destination_category or None, location=destination_location or None)
            else:
                add_device(device_name=destination_hostname or destination_ip, ip=destination_ip, model=destination_model,
                           category=destination_category or "radio", location=destination_location)

        result = add_link(source_ip=source_ip, destination_ip=destination_ip or None,
                          remark=remark, link_type=link_type, bandwidth="")
        if isinstance(result, Exception):
            flash(f"Error adding link: {result}")
        else:
            invalidate_cache_for("topology")
            flash("Link added")
        return redirect("/links")

    # ── GET: render links page ──────────────────────────────────────
    page         = request.args.get("page", 1, type=int)
    per_page     = request.args.get("per_page", 20, type=int)
    search_ip    = request.args.get("search_ip", "").strip()
    filter_status    = request.args.get("filter_status", "").strip()
    filter_category  = request.args.get("filter_category", "").strip()
    filter_location  = request.args.get("filter_location", "").strip()
    sort_by      = request.args.get("sort_by", "id").strip()
    sort_order   = request.args.get("sort_order", "desc").strip()

    if per_page not in (10, 20, 50, 100):
        per_page = 20
    page = max(1, page)

    valid_sort_cols = ("id", "source_ip", "destination_ip", "source_location",
                       "destination_location", "remark", "status")
    if sort_by not in valid_sort_cols:
        sort_by = "id"
    if sort_order not in ("asc", "desc"):
        sort_order = "desc"

    links, total, total_pages = get_links_paginated(
        page=page, per_page=per_page,
        search_ip=search_ip,
        filter_status=filter_status,
        filter_category=filter_category,
        filter_location=filter_location,
        sort_by=sort_by,
        sort_order=sort_order,
    )

    edit_link = None
    if edit_id:
        all_lks = get_all_links()
        edit_link = next((l for l in all_lks if l["id"] == edit_id), None)

    categories, locations = get_link_filter_options()

    need_location_ip = session.get("need_location_ip", None)
    pending_link = session.get("pending_link", None)

    return render_template(
        "links.html",
        links=links,
        edit_link=edit_link,
        active_page='links',
        page=page,
        per_page=per_page,
        total=total,
        total_pages=total_pages,
        search_ip=search_ip,
        filter_status=filter_status,
        filter_category=filter_category,
        filter_location=filter_location,
        sort_by=sort_by,
        sort_order=sort_order,
        categories=categories,
        locations=locations,
        need_location_ip=need_location_ip,
        pending_link=pending_link,
    )

# --- JSON API ---

@app.route("/api/alerts")
@login_required
def alerts_api():
    all_devices = all_devices_status()
    all_downs = [d for d in all_devices if not d["status"]]
    return jsonify(all_downs)


@app.route("/api/devices-list")
@login_required
def devices_list_api():
    return jsonify(all_devices_status())


@app.route("/api/devices-stats")
@login_required
def api_devices_stats():
    return jsonify(get_devices_stats())


@app.route("/api/devices-sidebar")
@login_required
def devices_sidebar_api():
    return jsonify(all_devices_status())


@app.route("/api/devices/delete/<ip>", methods=["DELETE"])
@login_required
@require_api_key
def delete_device_api(ip):
    result = delete_device(ip)
    if isinstance(result, Exception):
        return jsonify({"error": str(result)}), 500
    if result is False:
        return jsonify({"error": "Device not found"}), 404
    return jsonify({"status": "deleted", "ip": ip})



@app.route("/api/topology")
@login_required
def topology_api():
    cached = get_cached_topology("topology", ttl=30)
    if cached:
        return jsonify(cached)
    data = get_topology()
    cache_topology("topology", data)
    return jsonify(data)


@app.route("/api/topology/locations")
@login_required
def topology_locations_api():
    locations = get_locations()
    nodes = []
    for i, loc in enumerate(locations):
        nodes.append({
            "data": {
                "id": f"loc-{i}",
                "label": loc["name"],
                "status": loc["status"],
                "type": "location",
                "device_count": loc["device_count"],
                "down_count": loc["down_count"],
                "location_name": loc["name"]
            },
            "position": {"x": float(i * 400), "y": 0.0}
        })
    return jsonify({"nodes": nodes, "edges": []})


@app.route("/api/topology/location/<location_name>")
@login_required
def topology_location_api(location_name):
    devices = get_location_devices(location_name)
    links = get_location_links(location_name)
    nodes = []
    for i, d in enumerate(devices):
        nodes.append({
            "data": {
                "id": d["ip"],
                "label": f"{d['device_name']}\n{d['ip']}",
                "status": "GREEN" if d["status"] else "RED",
                "status_bool": d["status"],
                "location": location_name,
                "type": "device",
                "device_name": d["device_name"],
                "model": d["model"],
                "category": d["category"]
            },
            "position": {"x": float((i % 4) * 200), "y": float((i // 4) * 120 + 80)}
        })
    edges = []
    for link in links:
        edges.append({
            "data": {
                "id": f"link-{link['id']}",
                "source": link["source_ip"],
                "target": link["destination_ip"],
                "label": link["remark"] or "",
                "status": link["status"],
                "source_location": link["source_location"],
                "destination_location": link["destination_location"]
            }
        })
    return jsonify({"nodes": nodes, "edges": edges})


@app.route("/api/ping/now")
@login_required
def ping_now_api():
    from monitor.monitor import run as monitor_run
    ping_status["running"] = True
    def _run():
        try:
            result = monitor_run()
            ping_status["last_result"] = result
        except Exception as e:
            print(f"[PING] Manual ping failed: {e}")
            import traceback
            traceback.print_exc()
            ping_status["last_result"] = {"total": 0, "up": 0, "down": 0, "recovered": 0, "new_down": 0, "duration": 0}
        finally:
            ping_status["running"] = False
    threading.Thread(target=_run, daemon=True).start()
    return jsonify({"status": "started"})

@app.route("/api/ping/status")
@login_required
def ping_status_api():
    return jsonify(ping_status)


ping_status = {"running": False, "last_result": None}


@app.route("/api/monitor-status")
@login_required
def monitor_status_api():
    return jsonify(monitor_status)


@app.route("/api/health")
def health_api():
    db_ok = False
    try:
        s = SessionLocal()
        s.execute("SELECT 1")
        s.close()
        db_ok = True
    except Exception:
        pass

    scheduler_running = scheduler.running if hasattr(scheduler, "running") else False

    last_run = monitor_status.get("last_success")
    monitor_ok = False
    if last_run:
        from datetime import datetime as _dt
        try:
            age = (_dt.utcnow() - _dt.fromisoformat(last_run)).total_seconds()
            monitor_ok = age < 900
        except Exception:
            pass

    return jsonify({
        "database": db_ok,
        "scheduler": scheduler_running,
        "monitor": monitor_ok,
        "last_run": last_run,
        "monitor_status": monitor_status
    })


@app.route("/api/locations")
@login_required
def locations_api():
    return jsonify(get_locations())


@app.route("/api/dashboard-state", methods=["GET"])
@login_required
def get_dashboard_state_api():
    state = get_dashboard_state()
    return jsonify(state)


@app.route("/api/dashboard-state", methods=["POST"])
@login_required
@require_api_key
def save_dashboard_state_api():
    data = request.get_json(silent=True) or {}
    result = save_dashboard_state(
        zoom=data.get("zoom", 1.0),
        pan_x=data.get("pan_x", 0.0),
        pan_y=data.get("pan_y", 0.0),
        collapsed=data.get("collapsed", []),
        node_positions=data.get("node_positions", {}),
        layout_locked=data.get("layout_locked", False),
    )
    if isinstance(result, Exception):
        return jsonify({"error": str(result)}), 500
    state = get_dashboard_state()
    socketio.emit("dashboard_updated", state)
    return jsonify({"status": "saved"})


@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect("/")


@app.route("/admin/users", methods=["GET", "POST"])
@login_required
def admin_users():
    if current_user.user != "admin":
        flash("Access denied")
        return redirect("/dashboard")

    if request.method == "POST":
        action = request.form.get("action")
        if action == "add":
            username = request.form.get("username", "").strip()
            password = request.form.get("password", "").strip()
            must_change = request.form.get("must_change_password") == "on"
            if not username or not password:
                flash("Username and password are required")
            elif len(password) < 4:
                flash("Password must be at least 4 characters")
            else:
                ok, msg = create_user(username, password, must_change)
                flash(msg)
        elif action == "delete":
            user_id = request.form.get("user_id", type=int)
            if user_id and user_id != current_user.id:
                ok, msg = delete_user(user_id)
                flash(msg)
            elif user_id == current_user.id:
                flash("Cannot delete yourself")
        elif action == "reset_password":
            user_id = request.form.get("user_id", type=int)
            new_password = request.form.get("new_password", "").strip()
            if not new_password or len(new_password) < 4:
                flash("Password must be at least 4 characters")
            elif user_id:
                ok, msg = reset_user_password(user_id, new_password)
                flash(msg)
        return redirect("/admin/users")

    users = get_all_users()
    return render_template("admin_users.html", users=users)


@app.route("/topology-fullscreen")
@login_required
def topology_fullscreen():
    return render_template("topology_fullscreen.html")


@app.route("/api/updates-data")
@login_required
def updates_data_api():
    from datetime import timedelta
    s = SessionLocal()
    try:
        now = datetime.now(timezone.utc).replace(tzinfo=None)

        total_devices = s.query(Device).count()
        up_devices = s.query(Device).filter(Device.status == True).count()
        total_links = s.query(NetworkLink).count()
        up_links = s.query(NetworkLink).filter(NetworkLink.status == True).count()
        total_alerts = s.query(Alert).filter(Alert.resolved == False).count()
        total_users = s.query(User).count()

        cats = s.query(Category).all()
        categories = []
        for c in cats:
            cnt = s.query(Device).filter(Device.category_id == c.id).count()
            categories.append({"name": c.name, "count": cnt})

        recent_changes = []
        cutoff_24h = now - timedelta(hours=24)
        history = s.query(PingHistory).filter(
            PingHistory.ping_time >= cutoff_24h
        ).order_by(PingHistory.ping_time.desc()).limit(200).all()

        seen = {}
        for h in history:
            if h.device_ip not in seen:
                seen[h.device_ip] = {"ip": h.device_ip, "last_status": h.status, "last_time": h.ping_time.isoformat(), "latency": h.latency_ms}
            prev = seen[h.device_ip]
            if prev["last_status"] != h.status:
                recent_changes.append({
                    "ip": h.device_ip,
                    "from": "UP" if h.status else "DOWN",
                    "to": "DOWN" if h.status else "UP",
                    "time": h.ping_time.isoformat()
                })
                prev["last_status"] = h.status

        recent_changes = recent_changes[:20]

        devices_list = s.query(Device).order_by(Device.updated_at.desc().nullslast()).limit(50).all()
        device_updates = []
        for d in devices_list:
            device_updates.append({
                "ip": d.ip,
                "device_name": d.device_name,
                "hostname": d.hostname,
                "model": d.model,
                "location": d.location,
                "status": d.status,
                "last_seen": d.last_seen.isoformat() if d.last_seen else None,
                "created_at": d.created_at.isoformat() if d.created_at else None,
                "updated_at": d.updated_at.isoformat() if d.updated_at else None,
            })

        alerts = s.query(Alert).filter(Alert.resolved == False).order_by(Alert.created_at.desc()).limit(20).all()
        alert_list = []
        for a in alerts:
            alert_list.append({
                "id": a.id,
                "device_ip": a.device_ip,
                "created_at": a.created_at.isoformat() if a.created_at else None,
            })

        return jsonify({
            "summary": {
                "total_devices": total_devices,
                "up_devices": up_devices,
                "down_devices": total_devices - up_devices,
                "total_links": total_links,
                "up_links": up_links,
                "down_links": total_links - up_links,
                "active_alerts": total_alerts,
                "total_users": total_users,
                "categories": categories,
            },
            "recent_changes": recent_changes,
            "device_updates": device_updates,
            "alerts": alert_list,
        })
    finally:
        s.close()


if __name__ == "__main__":
    scheduler.start()
    Base.metadata.create_all(engine)
    create_default_admin()
    print("Database tables created.")
    debug_mode = os.getenv("DEBUG", "true").lower() == "true"
    ssl_cert = os.getenv("SSL_CERTFILE")
    ssl_key  = os.getenv("SSL_KEYFILE")

    if not debug_mode:
        app.config.update(SESSION_COOKIE_SECURE=True)

    ssl_context = None
    if ssl_cert and ssl_key:
        import os as _os
        if _os.path.isfile(ssl_cert) and _os.path.isfile(ssl_key):
            ssl_context = (ssl_cert, ssl_key)
            print(f"[SSL] TLS enabled: cert={ssl_cert} key={ssl_key}")
        else:
            print(f"[SSL] WARNING: cert or key file not found, falling back to plain HTTP")

    socketio.run(
        app,
        host=os.getenv("HOST", "0.0.0.0"),
        port=int(os.getenv("PORT", "5000")),
        debug=debug_mode,
        allow_unsafe_werkzeug=debug_mode,
        ssl_context=ssl_context
    )
